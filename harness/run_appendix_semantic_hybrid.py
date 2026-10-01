#!/usr/bin/env python3
"""Entropy-based and hybrid deletion experiments for the appendix datasets.

Supported datasets:
  - wiki_en  (Table 25)
  - reddit   (Table 26)
  - news_zh  (Table 29)
  - wiki_zh  (Table 30)
  - zhihu    (the published Zhihu rows of Table 31 come from run_zhihu_hybrid_experiments.py)

The script persists
  1. an aggregate results JSON (results_path, in harness/), which also holds the bucket calibration;
  2. per-sample JSONL records (samples_path, in samples/), one record per (method, rkeep, sample_idx).
Interrupted runs resume without recomputing finished samples. The shipped results files hold only the calibration
of the published runs, so a re-run reuses it. The published runs used Gemini 2.0 Flash (since retired); the
default decoder is now gemini-2.5-flash.
"""

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl
import torch
from bert_score import score as bert_score_fn
from transformers import AutoModelForCausalLM, AutoTokenizer, GPT2LMHeadModel, GPT2Tokenizer
from wordfreq import freq_to_zipf, get_frequency_dict, smash_numbers, zipf_frequency

sys.stdout.reconfigure(line_buffering=True)
sys.stderr.reconfigure(line_buffering=True)


ROOT = Path(__file__).parent
DATA_ROOT = ROOT.parent / "data"

RETENTION_RATES = [0.9, 0.7, 0.5, 0.3, 0.1]
ALPHA_VALUES = [0.3, 0.5, 0.7]
FREQ_HIGH = 5.0
FREQ_MID = 3.0
GOOGLE_API_KEY = os.environ["GOOGLE_API_KEY"]  # set your own key
ZH_FREQS = get_frequency_dict("zh", "best")

DATASETS = {
    "wiki_en": {
        "data_path": DATA_ROOT / "Wiki_text.xlsx",
        "sheet": "wf_3classes",
        "text_field": "wiki_text",
        "language": "en",
        "domain_label": "Wikipedia article excerpt",
        "bert_model": "distilroberta-base",
        "bert_lang": "en",
        "n_test": 100,
        "n_calib": 50,
        "results_path": ROOT / "wiki_hybrid_results_v2.json",
        "samples_path": Path(__file__).resolve().parents[1] / "samples" / "wiki_hybrid_samples.jsonl",
    },
    "reddit": {
        "data_path": DATA_ROOT / "Reddit_Dataset.xlsx",
        "sheet": "datasource",
        "text_field": "comment",
        "language": "en",
        "domain_label": "Reddit comment",
        "bert_model": "distilroberta-base",
        "bert_lang": "en",
        "n_test": 100,
        "n_calib": 50,
        "results_path": ROOT / "reddit_hybrid_results.json",
        "samples_path": Path(__file__).resolve().parents[1] / "samples" / "reddit_hybrid_samples.jsonl",
    },
    "news_zh": {
        "data_path": DATA_ROOT / "NEWS_Dataset.xlsx",
        "sheet": "Testing(Word Segmentation)",
        "text_field": "content",
        "token_field": "tok_flat_coarse",
        "language": "zh",
        "domain_label": "Chinese news article",
        "bert_model": "bert-base-chinese",
        "bert_lang": "zh",
        "surprisal_model": "uer/gpt2-chinese-cluecorpussmall",
        "n_test": 200,
        "n_calib": 50,
        "results_path": ROOT / "news_hybrid_results.json",
        "samples_path": Path(__file__).resolve().parents[1] / "samples" / "news_zh_hybrid_samples.jsonl",
    },
    "wiki_zh": {
        "data_path": DATA_ROOT / "Wiki_Dataset.xlsx",
        "sheet": "Testing(Word Segmentation)",
        "text_field": "content",
        "token_field": "tok_flat_coarse",
        "language": "zh",
        "domain_label": "Chinese Wikipedia article excerpt",
        "bert_model": "bert-base-chinese",
        "bert_lang": "zh",
        "surprisal_model": "uer/gpt2-chinese-cluecorpussmall",
        "n_test": 200,
        "n_calib": 50,
        "results_path": ROOT / "wiki_zh_hybrid_results.json",
        "samples_path": Path(__file__).resolve().parents[1] / "samples" / "wiki_zh_hybrid_samples.jsonl",
    },
    "zhihu": {
        "data_path": DATA_ROOT / "Zhihu_Dataset.xlsx",
        "sheet": "Testing(Word Segmentation)",
        "text_field": "RESPONSE",
        "token_field": "tok_flat_coarse",
        "language": "zh",
        "domain_label": "Zhihu answer",
        "bert_model": "bert-base-chinese",
        "bert_lang": "zh",
        "surprisal_model": "uer/gpt2-chinese-cluecorpussmall",
        "n_test": 200,
        "n_calib": 50,
        "results_path": ROOT / "zhihu_hybrid_results_v2.json",
        "samples_path": Path(__file__).resolve().parents[1] / "samples" / "zhihu_hybrid_samples.jsonl",
    },
}

import google.generativeai as genai

genai.configure(api_key=GOOGLE_API_KEY)
GEMINI_MODEL_NAME = "gemini-2.5-flash"  # default decoder; the published runs used gemini-2.0-flash, since retired
gemini = genai.GenerativeModel(GEMINI_MODEL_NAME)


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", choices=sorted(DATASETS.keys()), required=True)
    parser.add_argument("--n-test", type=int, default=None)
    parser.add_argument("--n-calib", type=int, default=None)
    parser.add_argument("--rates", type=float, nargs="+", default=RETENTION_RATES)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--use-cuda", action="store_true")
    return parser.parse_args()


def load_results(path: Path):
    if path.exists():
        return json.loads(path.read_text())
    return {}


def save_results(path: Path, data: dict):
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False))


def load_sample_records(path: Path):
    records = {}
    if not path.exists():
        return records
    with path.open() as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            key = (rec["method"], rec["rkeep"], int(rec["sample_idx"]))
            records[key] = rec
    return records


def append_sample_record(path: Path, record: dict):
    with path.open("a") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def split_segmented(segmented):
    if not segmented:
        return []
    return [tok for tok in str(segmented).split("/") if tok]


def normalize_compact_text(text):
    return re.sub(r"\s+", "", str(text))


def load_examples(cfg, n):
    wb = openpyxl.load_workbook(str(cfg["data_path"]), read_only=True)
    ws = wb[cfg["sheet"]]
    headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    header_to_idx = {str(name): idx for idx, name in enumerate(headers) if name is not None}

    examples = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=0):
        if cfg["language"] == "en":
            field = cfg["text_field"]
            field_idx = header_to_idx.get(field)
            if field_idx is None:
                field_idx = 0
            text = str(row[field_idx]).strip() if row[field_idx] else ""
            if len(text) < 100:
                continue
            examples.append(
                {
                    "id": row_idx,
                    "raw_text": text,
                    "compact_text": text,
                }
            )
        else:
            text = str(row[header_to_idx[cfg["text_field"]]]).strip() if row[header_to_idx[cfg["text_field"]]] else ""
            segmented = row[header_to_idx[cfg["token_field"]]] if cfg["token_field"] in header_to_idx else None
            tokens = split_segmented(segmented)
            compact_text = "".join(tokens) if tokens else normalize_compact_text(text)
            if not text or not compact_text or len(compact_text) < 80:
                continue
            examples.append(
                {
                    "id": row_idx,
                    "raw_text": text,
                    "compact_text": compact_text,
                    "tokens": tokens if tokens else list(compact_text),
                }
            )
        if len(examples) >= n:
            break
    wb.close()
    return examples


def call_gemini(prompt: str, max_retries: int = 5) -> str:
    for attempt in range(max_retries):
        try:
            response = gemini.generate_content(prompt)
            return response.text.strip()
        except Exception as e:
            print(f"  Gemini attempt {attempt + 1} failed: {e}", flush=True)
            if attempt < max_retries - 1:
                time.sleep(min(10, 2 + 2 * attempt))
            else:
                return ""


def reconstruct_prompt(cfg, degraded_text: str, original_len: int) -> str:
    if cfg["language"] == "zh":
        return (
            f"你是一个中文文本重建助手。输入是一段{cfg['domain_label']}，其中部分词语已被删除。"
            "请在不添加输入中没有依据的新事实、新人物或新事件的前提下，尽可能忠实地重建原文。"
            f"输出长度应大致接近原文（约{original_len}个字符）。"
            "只返回重建后的中文文本，不要加解释。\n\n"
            f"残缺文本：\n{degraded_text}"
        )
    return (
        f"You are a text reconstruction assistant. The input is a {cfg['domain_label']} "
        "from which some tokens have been deleted. Reconstruct the original text as faithfully as possible. "
        "Do not add new facts, names, or events not suggested by the input. "
        f"The output should be approximately the same length (~{original_len} characters). "
        "Return only the reconstructed text.\n\n"
        f"Degraded text:\n{degraded_text}"
    )


def load_surprisal_model(cfg, use_cuda: bool):
    if cfg["language"] == "zh":
        model_name = cfg["surprisal_model"]
        print(f"Loading surprisal model: {model_name}", flush=True)
        tokenizer = AutoTokenizer.from_pretrained(model_name)
        model = AutoModelForCausalLM.from_pretrained(model_name)
    else:
        print("Loading surprisal model: gpt2", flush=True)
        tokenizer = GPT2Tokenizer.from_pretrained("gpt2")
        model = GPT2LMHeadModel.from_pretrained("gpt2")
    model.eval()
    if use_cuda and torch.cuda.is_available():
        model = model.cuda()
    return tokenizer, model


def compute_token_surprisal_en(text, tokenizer, model):
    inputs = tokenizer(text, return_tensors="pt", truncation=True, max_length=1024)
    if next(model.parameters()).is_cuda:
        inputs = {k: v.cuda() for k, v in inputs.items()}
    with torch.no_grad():
        outputs = model(**inputs)
        logits = outputs.logits
    shift_logits = logits[:, :-1, :]
    shift_labels = inputs["input_ids"][:, 1:]
    log_probs = torch.nn.functional.log_softmax(shift_logits, dim=-1)
    token_log_probs = log_probs.gather(2, shift_labels.unsqueeze(-1)).squeeze(-1)
    surprisal = -token_log_probs.squeeze(0).cpu().numpy()
    median_s = float(np.median(surprisal)) if len(surprisal) else 0.0
    surprisal = np.concatenate([[median_s], surprisal]) if len(surprisal) else np.array([0.0])
    tokens = inputs["input_ids"].squeeze(0).cpu().tolist()
    token_strs = [tokenizer.decode([t]) for t in tokens]
    return token_strs, surprisal


def compute_char_surprisal_zh(text, tokenizer, model):
    encoded = tokenizer(
        text,
        return_tensors="pt",
        return_offsets_mapping=True,
        truncation=True,
        max_length=1024,
        add_special_tokens=False,
    )
    offsets = encoded.pop("offset_mapping")[0].tolist()
    if next(model.parameters()).is_cuda:
        encoded = {k: v.cuda() for k, v in encoded.items()}
    with torch.no_grad():
        outputs = model(**encoded)
        logits = outputs.logits
    shift_logits = logits[:, :-1, :]
    shift_labels = encoded["input_ids"][:, 1:]
    log_probs = torch.nn.functional.log_softmax(shift_logits, dim=-1)
    token_log_probs = log_probs.gather(2, shift_labels.unsqueeze(-1)).squeeze(-1)
    surprisal = -token_log_probs.squeeze(0).cpu().numpy()
    if len(surprisal) == 0:
        return np.array([0.0]), [(0, len(text))]
    median_s = float(np.median(surprisal))
    surprisal = np.concatenate([[median_s], surprisal])
    return surprisal, offsets


def compute_word_surprisal_zh(compact_text, words, tokenizer, model):
    token_surprisal, offsets = compute_char_surprisal_zh(compact_text, tokenizer, model)
    spans = []
    pos = 0
    for word in words:
        end = pos + len(word)
        spans.append((pos, end))
        pos = end
    if pos != len(compact_text):
        compact_text = "".join(words)
        token_surprisal, offsets = compute_char_surprisal_zh(compact_text, tokenizer, model)
        spans = []
        pos = 0
        for word in words:
            end = pos + len(word)
            spans.append((pos, end))
            pos = end
    word_scores = []
    for start, end in spans:
        vals = []
        for s, (tok_start, tok_end) in zip(token_surprisal, offsets):
            if tok_end <= start or tok_start >= end:
                continue
            vals.append(float(s))
        word_scores.append(float(np.mean(vals)) if vals else float(np.median(token_surprisal)))
    return np.array(word_scores, dtype=float)


def get_zipf_scores_en(token_strs):
    scores = []
    for tok in token_strs:
        word = tok.strip()
        if not word or not word[0].isalpha():
            scores.append(0.0)
        else:
            scores.append(zipf_frequency(word.lower(), "en"))
    return np.array(scores, dtype=float)


def get_zipf_scores_zh(words):
    scores = []
    for word in words:
        token = str(word).strip()
        if not token or re.fullmatch(r"\W+", token):
            scores.append(0.0)
            continue
        smashed = smash_numbers(token)
        freq = ZH_FREQS.get(smashed)
        if freq is None:
            char_freqs = [ZH_FREQS.get(ch) for ch in token if ZH_FREQS.get(ch) is not None]
            freq = float(np.mean(char_freqs)) if char_freqs else 1e-9
        scores.append(round(freq_to_zipf(freq), 2))
    return np.array(scores, dtype=float)


def assign_freq_bucket(zipf_scores):
    buckets = np.zeros(len(zipf_scores), dtype=int)
    buckets[zipf_scores >= FREQ_HIGH] = 2
    buckets[(zipf_scores >= FREQ_MID) & (zipf_scores < FREQ_HIGH)] = 1
    return buckets


def rank_normalize(arr):
    n = len(arr)
    if n <= 1:
        return np.zeros(n)
    order = np.argsort(arr, kind="stable")
    ranks = np.empty(n)
    ranks[order] = np.arange(n)
    return ranks / (n - 1)


def solve_lp(b_full, bucket_sizes, n_delete_total):
    cost = 1.0 - b_full
    order = np.argsort(cost)
    w = np.zeros(len(b_full))
    remaining = n_delete_total
    for k in order:
        cap = int(bucket_sizes[k])
        if cap <= 0:
            continue
        if remaining >= cap:
            w[k] = 1.0
            remaining -= cap
        else:
            w[k] = remaining / cap
            remaining = 0
            break
    return w


def evaluate_reconstructions(cfg, preds, refs):
    _, _, f1 = bert_score_fn(
        preds,
        refs,
        model_type=cfg["bert_model"],
        lang=cfg["bert_lang"],
        verbose=False,
        device="cpu",
    )
    return float(f1.mean().item()), float(f1.std().item())


def example_unit(cfg, ex, tokenizer, model):
    if cfg["language"] == "zh":
        words = ex["tokens"]
        surp = compute_word_surprisal_zh(ex["compact_text"], words, tokenizer, model)
        zipf_scores = get_zipf_scores_zh(words)
        return words, surp, zipf_scores
    token_strs, surp = compute_token_surprisal_en(ex["raw_text"], tokenizer, model)
    zipf_scores = get_zipf_scores_en(token_strs)
    return token_strs, surp, zipf_scores


def rebuild_text(cfg, units, keep_mask):
    kept = [u for u, keep in zip(units, keep_mask) if keep]
    return "".join(kept)


def calibrate_surprisal_buckets(cfg, calib_examples, tokenizer, model, results_path, all_results):
    if "calib_surprisal" in all_results:
        d = all_results["calib_surprisal"]
        print("  Surprisal calibration already exists, skipping.", flush=True)
        return np.array(d["b_full"]), d["thresholds"]
    print(f"\n--- Calibrating surprisal buckets (n={len(calib_examples)}) ---", flush=True)
    all_scores = []
    for ex in calib_examples:
        _, surp, _ = example_unit(cfg, ex, tokenizer, model)
        all_scores.extend(surp.tolist())
    t33 = float(np.percentile(all_scores, 33.33))
    t67 = float(np.percentile(all_scores, 66.67))
    print(f"  Surprisal thresholds: 33%={t33:.3f}, 67%={t67:.3f}", flush=True)

    bucket_preds = [[], [], []]
    bucket_refs = [[], [], []]
    for i, ex in enumerate(calib_examples):
        units, surp, _ = example_unit(cfg, ex, tokenizer, model)
        for bk in range(3):
            if bk == 0:
                delete_mask = surp <= t33
            elif bk == 1:
                delete_mask = (surp > t33) & (surp <= t67)
            else:
                delete_mask = surp > t67
            degraded = rebuild_text(cfg, units, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(cfg, degraded, len(ex["raw_text"])))
            bucket_preds[bk].append(rec)
            bucket_refs[bk].append(ex["raw_text"])
            time.sleep(0.1)
        if (i + 1) % 10 == 0:
            print(f"  calibration {i + 1}/{len(calib_examples)}", flush=True)

    b_full = np.zeros(3)
    for bk in range(3):
        mean_f1, _ = evaluate_reconstructions(cfg, bucket_preds[bk], bucket_refs[bk])
        b_full[bk] = mean_f1
        print(f"  Surprisal bucket {bk}: B_k^full={b_full[bk]:.4f}", flush=True)
    all_results["calib_surprisal"] = {"b_full": b_full.tolist(), "thresholds": [t33, t67]}
    save_results(results_path, all_results)
    return b_full, [t33, t67]


def calibrate_freq_buckets(cfg, calib_examples, tokenizer, model, results_path, all_results):
    if "calib_frequency" in all_results:
        d = all_results["calib_frequency"]
        print("  Frequency calibration already exists, skipping.", flush=True)
        return np.array(d["b_full"])
    print(f"\n--- Calibrating frequency buckets (n={len(calib_examples)}) ---", flush=True)
    bucket_preds = [[], [], []]
    bucket_refs = [[], [], []]
    for i, ex in enumerate(calib_examples):
        units, _, zipf_scores = example_unit(cfg, ex, tokenizer, model)
        bucket_ids = assign_freq_bucket(zipf_scores)
        for bk in range(3):
            delete_mask = bucket_ids == bk
            degraded = rebuild_text(cfg, units, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(cfg, degraded, len(ex["raw_text"])))
            bucket_preds[bk].append(rec)
            bucket_refs[bk].append(ex["raw_text"])
            time.sleep(0.1)
        if (i + 1) % 10 == 0:
            print(f"  calibration {i + 1}/{len(calib_examples)}", flush=True)

    b_full = np.zeros(3)
    for bk in range(3):
        mean_f1, _ = evaluate_reconstructions(cfg, bucket_preds[bk], bucket_refs[bk])
        b_full[bk] = mean_f1
        print(f"  Frequency bucket {bk}: B_k^full={b_full[bk]:.4f}", flush=True)
    all_results["calib_frequency"] = {"b_full": b_full.tolist()}
    save_results(results_path, all_results)
    return b_full


def sample_record_base(cfg, method_key, rkeep, ex, sample_idx, degraded, reconstruction):
    return {
        "dataset": cfg["results_path"].stem.replace("_hybrid_results", ""),
        "language": cfg["language"],
        "decoder_model": GEMINI_MODEL_NAME,
        "method": method_key,
        "rkeep": str(rkeep),
        "sample_idx": int(sample_idx),
        "source_row_id": int(ex["id"]),
        "reference": ex["raw_text"],
        "degraded": degraded,
        "reconstruction": reconstruction,
    }


def run_method(
    cfg,
    method_key,
    display_name,
    examples,
    tokenizer,
    model,
    results_path,
    samples_path,
    all_results,
    sample_records,
    rates,
    builder_fn,
):
    results = all_results.get(method_key, {})
    for rkeep in rates:
        rate_key = str(rkeep)
        print(f"\n=== {display_name} rkeep={rkeep} ===", flush=True)
        preds = []
        refs = []
        complete = 0
        for i, ex in enumerate(examples):
            existing = sample_records.get((method_key, rate_key, i))
            if existing is not None:
                preds.append(existing["reconstruction"])
                refs.append(existing["reference"])
                complete += 1
                continue
            degraded = builder_fn(ex, rkeep)
            rec = call_gemini(reconstruct_prompt(cfg, degraded, len(ex["raw_text"])))
            record = sample_record_base(cfg, method_key, rkeep, ex, i, degraded, rec)
            append_sample_record(samples_path, record)
            sample_records[(method_key, rate_key, i)] = record
            preds.append(rec)
            refs.append(ex["raw_text"])
            complete += 1
            if complete % 10 == 0:
                print(f"  {complete}/{len(examples)} done", flush=True)
            time.sleep(0.1)

        mean_f1, std_f1 = evaluate_reconstructions(cfg, preds, refs)
        print(f"  BERTScore F1: {mean_f1:.4f} +/- {std_f1:.4f}", flush=True)
        results[rate_key] = {"mean_f1": round(mean_f1, 4), "std_f1": round(std_f1, 4), "n": len(examples)}
        all_results[method_key] = results
        save_results(results_path, all_results)
    return results


def print_summary(all_results, rates):
    cols = [
        ("Entropy", all_results.get("entropy", {})),
        ("Entropy-LP", all_results.get("entropy_lp", {})),
        ("Ent.-FreqBkt", all_results.get("entropy_in_freqbuckets", {})),
        ("Hybrid-0.3", all_results.get("hybrid_alpha_0.3", {})),
        ("Hybrid-0.5", all_results.get("hybrid_alpha_0.5", {})),
        ("Hybrid-0.7", all_results.get("hybrid_alpha_0.7", {})),
    ]
    print("\n" + "=" * 90, flush=True)
    print("SUMMARY", flush=True)
    print("=" * 90, flush=True)
    header = f"{'rkeep':<8}" + "".join(f"{label:<16}" for label, _ in cols)
    print(header, flush=True)
    print("-" * len(header), flush=True)
    for rkeep in rates:
        row = f"{rkeep:<8}"
        for _, data in cols:
            v = data.get(str(rkeep), {})
            row += f"{str(v.get('mean_f1', '---')):<16}"
        print(row, flush=True)


def main():
    args = parse_args()
    cfg = dict(DATASETS[args.dataset])
    if args.n_test is not None:
        cfg["n_test"] = args.n_test
    if args.n_calib is not None:
        cfg["n_calib"] = args.n_calib

    print(f"Dataset: {args.dataset}", flush=True)
    print(f"Results: {cfg['results_path']}", flush=True)
    print(f"Samples: {cfg['samples_path']}", flush=True)

    examples = load_examples(cfg, cfg["n_test"])
    calib_examples = examples[: cfg["n_calib"]]
    print(f"Loaded {len(examples)} examples", flush=True)

    tokenizer, model = load_surprisal_model(cfg, args.use_cuda)
    if args.dry_run:
        ex = examples[0]
        units, surp, zipf_scores = example_unit(cfg, ex, tokenizer, model)
        print(f"Dry run ok: units={len(units)} surprisal={len(surp)} zipf={len(zipf_scores)}", flush=True)
        return

    all_results = load_results(cfg["results_path"])
    all_results.setdefault(
        "_meta",
        {
            "dataset": args.dataset,
            "decoder_model": GEMINI_MODEL_NAME,
            "bert_model": cfg["bert_model"],
            "n_test": len(examples),
            "n_calib": len(calib_examples),
            "rates": args.rates,
        },
    )
    sample_records = load_sample_records(cfg["samples_path"])

    b_full_surp, thresholds = calibrate_surprisal_buckets(
        cfg, calib_examples, tokenizer, model, cfg["results_path"], all_results
    )
    b_full_freq = calibrate_freq_buckets(
        cfg, calib_examples, tokenizer, model, cfg["results_path"], all_results
    )

    def build_entropy(ex, rkeep):
        units, surp, _ = example_unit(cfg, ex, tokenizer, model)
        n_keep = max(1, int(len(units) * rkeep))
        keep_idx = set(np.argsort(surp)[-n_keep:])
        keep_mask = np.array([j in keep_idx for j in range(len(units))], dtype=bool)
        return rebuild_text(cfg, units, keep_mask)

    t33, t67 = thresholds

    def build_entropy_lp(ex, rkeep):
        units, surp, _ = example_unit(cfg, ex, tokenizer, model)
        n = len(units)
        n_del = max(0, int(n * (1 - rkeep)))
        bucket_ids = np.zeros(n, dtype=int)
        bucket_ids[(surp > t33) & (surp <= t67)] = 1
        bucket_ids[surp > t67] = 2
        bucket_sizes = np.array([np.sum(bucket_ids == k) for k in range(3)])
        w = solve_lp(b_full_surp, bucket_sizes, n_del)
        delete_mask = np.zeros(n, dtype=bool)
        for bk in range(3):
            idx = np.where(bucket_ids == bk)[0]
            n_del_bk = int(round(w[bk] * len(idx)))
            if n_del_bk > 0:
                cheapest = idx[np.argsort(surp[idx])][:n_del_bk]
                delete_mask[cheapest] = True
        return rebuild_text(cfg, units, ~delete_mask)

    def build_entropy_in_freqbuckets(ex, rkeep):
        units, surp, zipf_scores = example_unit(cfg, ex, tokenizer, model)
        n = len(units)
        n_del = max(0, int(n * (1 - rkeep)))
        bucket_ids = assign_freq_bucket(zipf_scores)
        bucket_sizes = np.array([np.sum(bucket_ids == k) for k in range(3)])
        w = solve_lp(b_full_freq, bucket_sizes, n_del)
        delete_mask = np.zeros(n, dtype=bool)
        for bk in range(3):
            idx = np.where(bucket_ids == bk)[0]
            n_del_bk = int(round(w[bk] * len(idx)))
            if n_del_bk > 0:
                cheapest = idx[np.argsort(surp[idx])][:n_del_bk]
                delete_mask[cheapest] = True
        return rebuild_text(cfg, units, ~delete_mask)

    run_method(
        cfg,
        "entropy",
        "Entropy",
        examples,
        tokenizer,
        model,
        cfg["results_path"],
        cfg["samples_path"],
        all_results,
        sample_records,
        args.rates,
        build_entropy,
    )
    run_method(
        cfg,
        "entropy_lp",
        "Entropy-LP",
        examples,
        tokenizer,
        model,
        cfg["results_path"],
        cfg["samples_path"],
        all_results,
        sample_records,
        args.rates,
        build_entropy_lp,
    )
    run_method(
        cfg,
        "entropy_in_freqbuckets",
        "Entropy-in-FreqBuckets",
        examples,
        tokenizer,
        model,
        cfg["results_path"],
        cfg["samples_path"],
        all_results,
        sample_records,
        args.rates,
        build_entropy_in_freqbuckets,
    )

    for alpha in ALPHA_VALUES:
        def build_hybrid(ex, rkeep, alpha=alpha):
            units, surp, zipf_scores = example_unit(cfg, ex, tokenizer, model)
            n_keep = max(1, int(len(units) * rkeep))
            rank_freq = rank_normalize(zipf_scores)
            inv_surp = np.where(surp > 0, 1.0 / surp, 0.0)
            rank_pred = rank_normalize(inv_surp)
            combined = alpha * rank_freq + (1.0 - alpha) * rank_pred
            keep_indices = set(np.argsort(combined)[:n_keep])
            keep_mask = np.array([j in keep_indices for j in range(len(units))], dtype=bool)
            return rebuild_text(cfg, units, keep_mask)

        run_method(
            cfg,
            f"hybrid_alpha_{alpha}",
            f"Hybrid-{alpha}",
            examples,
            tokenizer,
            model,
            cfg["results_path"],
            cfg["samples_path"],
            all_results,
            sample_records,
            args.rates,
            build_hybrid,
        )

    print_summary(all_results, args.rates)


if __name__ == "__main__":
    main()
