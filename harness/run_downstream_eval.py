#!/usr/bin/env python3
"""
Downstream task evaluation for skeleton-only and reconstruct-then-task.

This script is intentionally lightweight and self-contained (similar style to the other `run_*.py` files).

Tasks (BBC spreadsheet only):
  1) Topic classification over the BBC categories: supervised TF-IDF + linear SVM (default sklearn backend) or
     zero-shot NLI (hf backend).
  2) Retrieval proxy: use article summaries as queries and retrieve matching articles (Recall@K).

Evaluation tracks:
  - original (upper bound)
  - skeleton-only
  - reconstructed (decoder applied to skeleton)

Notes
-----
Some components are optional and lazily imported:
  - Compression methods that require GPT-2 / wordfreq need `torch`, `transformers`, `wordfreq`.
  - "hf" backend needs `transformers` (topic) and `sentence-transformers` (retrieval).
  - Default "sklearn" backend needs `scikit-learn` only.
  - Reconstruction can use either Gemini (google-generativeai) or a local HF causal LM.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import random
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

sys.stdout.reconfigure(line_buffering=True)
sys.stderr.reconfigure(line_buffering=True)

import numpy as np
import openpyxl


DATA_PATH = Path(__file__).parent.parent / "data" / "BBC_news_latest.xlsx"
RESULTS_PATH = Path(__file__).parent / "downstream_eval_results.json"


# ----------------------------
# Data
# ----------------------------
@dataclass(frozen=True)
class BBCExample:
    row_id: int
    category: str
    article: str
    summary: str


def load_bbc_examples(n: int, min_chars: int = 200) -> list[BBCExample]:
    wb = openpyxl.load_workbook(str(DATA_PATH), read_only=True)
    ws = wb["datasource"]
    out: list[BBCExample] = []
    # header: index, Category, Articles, Summaries
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 4:
            continue
        row_id = int(row[0]) if row[0] is not None else len(out)
        category = str(row[1]).strip() if row[1] else ""
        article = str(row[2]) if row[2] else ""
        summary = str(row[3]) if row[3] else ""
        if not category or len(article) < min_chars:
            continue
        out.append(BBCExample(row_id=row_id, category=category, article=article, summary=summary))
        if len(out) >= n:
            break
    wb.close()
    return out


def load_bbc_pool(max_rows: int | None = None, min_chars: int = 200) -> list[BBCExample]:
    wb = openpyxl.load_workbook(str(DATA_PATH), read_only=True)
    ws = wb["datasource"]
    out: list[BBCExample] = []
    # header: index, Category, Articles, Summaries
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 4:
            continue
        row_id = int(row[0]) if row[0] is not None else len(out)
        category = str(row[1]).strip() if row[1] else ""
        article = str(row[2]) if row[2] else ""
        summary = str(row[3]) if row[3] else ""
        if not category or len(article) < min_chars:
            continue
        out.append(BBCExample(row_id=row_id, category=category, article=article, summary=summary))
        if max_rows is not None and len(out) >= max_rows:
            break
    wb.close()
    return out


# ----------------------------
# Compression (skeleton)
# ----------------------------
def truncation_chars(text: str, rkeep: float) -> str:
    n = len(text)
    n_keep = max(1, int(n * rkeep))
    return text[:n_keep]


def random_char_deletion(text: str, rkeep: float, seed: int = 0) -> str:
    rng = random.Random(seed)
    n = len(text)
    n_keep = max(1, int(n * rkeep))
    if n_keep >= n:
        return text
    keep = set(rng.sample(range(n), n_keep))
    return "".join(ch for i, ch in enumerate(text) if i in keep)


def step_char_deletion(text: str, rkeep: float) -> str:
    n = len(text)
    n_keep = max(1, int(n * rkeep))
    if n_keep >= n:
        return text
    keep_idx = set(np.linspace(0, n - 1, n_keep, dtype=int).tolist())
    return "".join(ch for i, ch in enumerate(text) if i in keep_idx)


def _choose_evenly_spaced_indices(indices: list[int], k: int) -> list[int]:
    if k <= 0 or not indices:
        return []
    if k >= len(indices):
        return list(indices)
    pick_idx = np.linspace(0, len(indices) - 1, k, dtype=int).tolist()
    return [indices[i] for i in pick_idx]


def _apply_stage_deletions(candidates: list[int], remove_set: set[int], target_remove: int):
    remaining = [idx for idx in candidates if idx not in remove_set]
    need = target_remove - len(remove_set)
    if need <= 0 or not remaining:
        return
    for idx in _choose_evenly_spaced_indices(remaining, min(need, len(remaining))):
        remove_set.add(idx)


def wordlen_char_deletion(text: str, rkeep: float) -> str:
    """
    Pure-python approximation of the paper's WordLen heuristic.
    Deletes in stages: whitespace -> vowels (non-initial) -> short words -> long tail -> punct/digits.
    """
    import re

    n = len(text)
    n_keep = max(1, int(n * rkeep))
    target_remove = max(0, n - n_keep)
    if target_remove == 0:
        return text

    remove_set: set[int] = set()
    token_pattern = re.compile(r"\s+|[A-Za-z]+(?:'[A-Za-z]+)?|\d+|[^\w\s]+")
    vowels = set("aeiouAEIOU")

    whitespace_positions: list[int] = []
    vowel_positions: list[int] = []
    short_word_positions: list[int] = []
    long_word_tail_positions: list[int] = []
    punct_digit_positions: list[int] = []

    for match in token_pattern.finditer(text):
        tok = match.group(0)
        start, end = match.span()
        positions = list(range(start, end))
        if tok.isspace():
            whitespace_positions.extend(positions)
        elif tok.isalpha():
            if len(tok) >= 3:
                for off, ch in enumerate(tok):
                    if off > 0 and ch in vowels:
                        vowel_positions.append(start + off)
            if len(tok) <= 2:
                short_word_positions.extend(positions)
            if len(tok) > 6:
                long_word_tail_positions.extend(range(start + 4, end))
        else:
            punct_digit_positions.extend(positions)

    _apply_stage_deletions(whitespace_positions, remove_set, target_remove)
    _apply_stage_deletions(vowel_positions, remove_set, target_remove)
    _apply_stage_deletions(short_word_positions, remove_set, target_remove)
    _apply_stage_deletions(long_word_tail_positions, remove_set, target_remove)
    _apply_stage_deletions(punct_digit_positions, remove_set, target_remove)

    if len(remove_set) < target_remove:
        remaining = [idx for idx in range(n) if idx not in remove_set]
        for idx in _choose_evenly_spaced_indices(remaining, target_remove - len(remove_set)):
            remove_set.add(idx)

    return "".join(ch for i, ch in enumerate(text) if i not in remove_set)


def _lazy_import_transformers():
    try:
        import torch  # noqa: F401
        from transformers import GPT2LMHeadModel, GPT2Tokenizer  # noqa: F401
    except Exception as e:  # pragma: no cover
        raise RuntimeError(
            "Missing deps for GPT-2-based compression. Install `torch` and `transformers`."
        ) from e


def entropy_token_deletion(text: str, rkeep: float) -> str:
    """
    Delete low-surprisal GPT-2 tokens (keeps highest-surprisal tokens).
    """
    _lazy_import_transformers()
    import torch
    from transformers import GPT2LMHeadModel, GPT2Tokenizer

    tokenizer = GPT2Tokenizer.from_pretrained("gpt2")
    model = GPT2LMHeadModel.from_pretrained("gpt2")
    model.eval()
    if torch.cuda.is_available():
        model = model.cuda()

    inputs = tokenizer(text, return_tensors="pt", truncation=True, max_length=1024)
    if torch.cuda.is_available():
        inputs = {k: v.cuda() for k, v in inputs.items()}
    with torch.no_grad():
        logits = model(**inputs).logits
    shift_logits = logits[:, :-1, :]
    shift_labels = inputs["input_ids"][:, 1:]
    log_probs = torch.nn.functional.log_softmax(shift_logits, dim=-1)
    token_log_probs = log_probs.gather(2, shift_labels.unsqueeze(-1)).squeeze(-1)
    surprisal = (-token_log_probs.squeeze(0)).cpu().numpy()
    median_s = float(np.median(surprisal))
    surprisal = np.concatenate([[median_s], surprisal])

    tokens = inputs["input_ids"].squeeze(0).cpu().tolist()
    token_strs = [tokenizer.decode([t]) for t in tokens]
    n = len(token_strs)
    n_keep = max(1, int(n * rkeep))
    keep_indices = set(np.argsort(surprisal)[-n_keep:])
    return "".join(tok for i, tok in enumerate(token_strs) if i in keep_indices)


def wordfreq_token_deletion(text: str, rkeep: float) -> str:
    """
    Delete highest-frequency GPT-2 tokens (keeps lowest Zipf frequency tokens).
    """
    _lazy_import_transformers()
    try:
        from wordfreq import zipf_frequency
    except Exception as e:  # pragma: no cover
        raise RuntimeError("Missing dep `wordfreq` for wordfreq deletion.") from e

    import torch
    from transformers import GPT2LMHeadModel, GPT2Tokenizer

    tokenizer = GPT2Tokenizer.from_pretrained("gpt2")
    model = GPT2LMHeadModel.from_pretrained("gpt2")
    model.eval()
    if torch.cuda.is_available():
        model = model.cuda()

    inputs = tokenizer(text, return_tensors="pt", truncation=True, max_length=1024)
    tokens = inputs["input_ids"].squeeze(0).cpu().tolist()
    token_strs = [tokenizer.decode([t]) for t in tokens]
    n = len(token_strs)
    n_keep = max(1, int(n * rkeep))

    zipf_scores = []
    for tok in token_strs:
        word = tok.strip()
        if not word or not word[0].isalpha():
            zipf_scores.append(0.0)
        else:
            zipf_scores.append(zipf_frequency(word.lower(), "en"))
    zipf_scores = np.array(zipf_scores, dtype=float)
    keep_indices = set(np.argsort(zipf_scores)[:n_keep])
    return "".join(tok for i, tok in enumerate(token_strs) if i in keep_indices)


def get_compressor(name: str, seed: int) -> Callable[[str, float], str]:
    name = name.strip().lower()
    if name == "trunc":
        return truncation_chars
    if name == "random":
        return lambda t, rk: random_char_deletion(t, rk, seed=seed)
    if name == "step":
        return step_char_deletion
    if name == "wordlen":
        return wordlen_char_deletion
    if name == "entropy":
        return entropy_token_deletion
    if name == "wordfreq":
        return wordfreq_token_deletion
    raise ValueError(f"Unknown compressor: {name}")


# ----------------------------
# Reconstruction
# ----------------------------
def reconstruct_prompt(degraded: str, original_len: int) -> str:
    return (
        "You are a text reconstruction assistant. The input is an English news article "
        "from which some characters/tokens have been deleted. Your task is to reconstruct "
        "the original article as faithfully as possible. Do not add new facts, names, or "
        "events not suggested by the input. The output should be approximately the same "
        f"length as the original (~{original_len} characters). Return only the reconstructed text.\n\n"
        f"Degraded text:\n{degraded}"
    )


class Reconstructor:
    def reconstruct(self, degraded: str, original_len: int) -> str:
        raise NotImplementedError


class GeminiReconstructor(Reconstructor):
    def __init__(self, model_name: str, api_key: str | None, max_retries: int = 3):
        try:
            import google.generativeai as genai
        except Exception as e:  # pragma: no cover
            raise RuntimeError("Missing dep `google-generativeai` for Gemini reconstruction.") from e

        api_key = api_key or os.environ.get("GOOGLE_API_KEY")
        if not api_key:
            raise RuntimeError("Gemini decoder selected but no API key found (set `GOOGLE_API_KEY`).")

        genai.configure(api_key=api_key)
        self._model = genai.GenerativeModel(model_name)
        self._max_retries = max_retries

    def reconstruct(self, degraded: str, original_len: int) -> str:
        prompt = reconstruct_prompt(degraded, original_len)
        for attempt in range(self._max_retries):
            try:
                resp = self._model.generate_content(prompt)
                return (resp.text or "").strip()
            except Exception as e:
                print(f"  Gemini attempt {attempt+1} failed: {e}", flush=True)
                if attempt < self._max_retries - 1:
                    time.sleep(2)
        return ""


class HFReconstructor(Reconstructor):
    def __init__(self, model_name: str, max_new_tokens: int = 512):
        try:
            import torch
            from transformers import AutoModelForCausalLM, AutoTokenizer
        except Exception as e:  # pragma: no cover
            raise RuntimeError("Missing deps for HF reconstruction. Install `torch` and `transformers`.") from e

        self._torch = torch
        self._tok = AutoTokenizer.from_pretrained(model_name, use_fast=True)
        self._model = AutoModelForCausalLM.from_pretrained(model_name)
        self._model.eval()
        if torch.cuda.is_available():
            self._model = self._model.cuda()
        self._max_new_tokens = max_new_tokens

    def reconstruct(self, degraded: str, original_len: int) -> str:
        prompt = reconstruct_prompt(degraded, original_len)
        inputs = self._tok(prompt, return_tensors="pt", truncation=True, max_length=2048)
        if self._torch.cuda.is_available():
            inputs = {k: v.cuda() for k, v in inputs.items()}
        with self._torch.no_grad():
            out = self._model.generate(
                **inputs,
                max_new_tokens=self._max_new_tokens,
                do_sample=False,
                eos_token_id=self._tok.eos_token_id,
            )
        text = self._tok.decode(out[0], skip_special_tokens=True)
        # Attempt to strip the prompt prefix if the model echoed it.
        marker = "Degraded text:"
        if marker in text:
            text = text.split(marker, 1)[-1].strip()
        return text.strip()


def build_reconstructor(kind: str, *, gemini_model: str, hf_model: str, api_key: str | None) -> Reconstructor | None:
    kind = kind.strip().lower()
    if kind in ("none", "off", ""):
        return None
    if kind == "gemini":
        return GeminiReconstructor(model_name=gemini_model, api_key=api_key)
    if kind == "hf":
        return HFReconstructor(model_name=hf_model)
    raise ValueError(f"Unknown decoder: {kind}")


# ----------------------------
# Metrics
# ----------------------------
def accuracy(y_true: list[str], y_pred: list[str]) -> float:
    if not y_true:
        return 0.0
    return sum(int(a == b) for a, b in zip(y_true, y_pred)) / len(y_true)


def macro_f1(y_true: list[str], y_pred: list[str], labels: list[str]) -> float:
    if not y_true:
        return 0.0
    # per-class counts
    tp = {l: 0 for l in labels}
    fp = {l: 0 for l in labels}
    fn = {l: 0 for l in labels}
    for t, p in zip(y_true, y_pred):
        if p == t:
            tp[t] += 1
        else:
            fp[p] += 1
            fn[t] += 1

    f1s = []
    for l in labels:
        prec = tp[l] / (tp[l] + fp[l]) if (tp[l] + fp[l]) > 0 else 0.0
        rec = tp[l] / (tp[l] + fn[l]) if (tp[l] + fn[l]) > 0 else 0.0
        f1 = (2 * prec * rec / (prec + rec)) if (prec + rec) > 0 else 0.0
        f1s.append(f1)
    return float(sum(f1s) / len(f1s)) if f1s else 0.0


def recall_at_k(sim: np.ndarray, k: int) -> float:
    """
    sim: (n_queries, n_docs), larger is better. Assumes correct doc is diagonal.
    """
    n = sim.shape[0]
    hits = 0
    for i in range(n):
        topk = np.argpartition(-sim[i], kth=min(k, n - 1))[:k]
        if i in topk:
            hits += 1
    return hits / n if n > 0 else 0.0


# ----------------------------
# Downstream tasks
# ----------------------------
def _lazy_import_sklearn():
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer  # noqa: F401
        from sklearn.svm import LinearSVC  # noqa: F401
    except Exception as e:  # pragma: no cover
        raise RuntimeError("Missing dep `scikit-learn` for sklearn backend.") from e


def _lazy_import_transformers_pipeline():
    try:
        from transformers import pipeline  # noqa: F401
    except Exception as e:  # pragma: no cover
        raise RuntimeError("Missing dep `transformers` for downstream tasks.") from e


def topic_classify_zero_shot(
    texts: list[str],
    labels: list[str],
    model_name: str,
    batch_size: int = 8,
) -> list[str]:
    _lazy_import_transformers_pipeline()
    from transformers import pipeline

    clf = pipeline("zero-shot-classification", model=model_name)
    preds: list[str] = []
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        out = clf(batch, candidate_labels=labels)
        # `out` is list[dict] when batch input; dict when single input.
        if isinstance(out, dict):
            out = [out]
        for r in out:
            preds.append(r["labels"][0])
    return preds


def _lazy_import_sentence_transformers():
    try:
        from sentence_transformers import SentenceTransformer  # noqa: F401
    except Exception as e:  # pragma: no cover
        raise RuntimeError("Missing dep `sentence-transformers` for retrieval task.") from e


def embed_texts(texts: list[str], model_name: str, batch_size: int = 32) -> np.ndarray:
    _lazy_import_sentence_transformers()
    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(model_name)
    emb = model.encode(texts, batch_size=batch_size, normalize_embeddings=True, show_progress_bar=False)
    return np.asarray(emb, dtype=np.float32)


def build_sklearn_models(
    train_texts: list[str],
    train_labels: list[str],
    *,
    max_features: int,
    ngram_max: int,
) -> tuple[object, object]:
    _lazy_import_sklearn()
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.svm import LinearSVC

    vectorizer = TfidfVectorizer(
        stop_words="english",
        max_features=max_features if max_features > 0 else None,
        ngram_range=(1, max(1, ngram_max)),
        norm="l2",
    )
    x_train = vectorizer.fit_transform(train_texts)
    clf = LinearSVC()
    clf.fit(x_train, train_labels)
    return vectorizer, clf


def sklearn_predict(vectorizer: object, clf: object, texts: list[str]) -> list[str]:
    x = vectorizer.transform(texts)
    return clf.predict(x).tolist()


def sklearn_retrieval_sim(vectorizer: object, query_texts: list[str], doc_texts: list[str]) -> np.ndarray:
    q = vectorizer.transform(query_texts)
    d = vectorizer.transform(doc_texts)
    sim = (q @ d.T).toarray()
    return np.asarray(sim, dtype=np.float32)


# ----------------------------
# Main runner
# ----------------------------
def _cap_text(s: str, max_chars: int) -> str:
    s = s.strip()
    if max_chars <= 0 or len(s) <= max_chars:
        return s
    return s[:max_chars]


def run(args: argparse.Namespace) -> dict:
    pool = load_bbc_pool(max_rows=args.max_rows, min_chars=args.min_article_chars)
    if len(pool) < max(1, args.n_test):
        raise RuntimeError("No BBC examples loaded. Check `data/BBC_news_latest.xlsx` and filters.")

    rng = random.Random(args.seed)
    rng.shuffle(pool)
    train_examples = pool[: args.n_train] if args.n_train > 0 else []
    test_examples = pool[args.n_train : args.n_train + args.n_test]
    if not test_examples:
        raise RuntimeError("No test examples selected (check --n-train/--n-test/--max-rows).")

    labels = sorted({ex.category for ex in test_examples + train_examples})
    compressor = get_compressor(args.compressor, seed=args.seed)
    reconstructor = build_reconstructor(
        args.decoder,
        gemini_model=args.gemini_model,
        hf_model=args.hf_model,
        api_key=args.gemini_api_key,
    )

    # Prepare texts
    originals = [_cap_text(ex.article, args.max_article_chars) for ex in test_examples]
    summaries = [_cap_text(ex.summary, args.max_summary_chars) for ex in test_examples]
    y_true = [ex.category for ex in test_examples]
    train_originals = [_cap_text(ex.article, args.max_article_chars) for ex in train_examples]
    y_train = [ex.category for ex in train_examples]

    results: dict = {
        "dataset": "BBC_news_latest.xlsx",
        "backend": args.backend,
        "n_train": len(train_examples),
        "n_test": len(test_examples),
        "labels": labels,
        "compressor": args.compressor,
        "decoder": args.decoder,
        "rkeeps": args.rkeeps,
        "tasks": ["topic", "summary_to_article_retrieval"],
        "topic_method": "zero-shot NLI" if args.backend == "hf" else "supervised TF-IDF + linear SVM",
        "runs": {},
    }

    sklearn_vectorizer = None
    sklearn_clf = None
    if args.backend == "sklearn":
        if not train_examples:
            raise RuntimeError("sklearn backend requires --n-train > 0 to fit the classifier/vectorizer.")
        sklearn_vectorizer, sklearn_clf = build_sklearn_models(
            train_originals,
            y_train,
            max_features=args.tfidf_max_features,
            ngram_max=args.tfidf_ngram_max,
        )

    def eval_variant(name: str, texts: list[str]) -> dict:
        out: dict = {}
        if "topic" in args.tasks:
            if args.backend == "hf":
                y_pred = topic_classify_zero_shot(
                    texts, labels, model_name=args.topic_model, batch_size=args.topic_batch_size
                )
            else:
                y_pred = sklearn_predict(sklearn_vectorizer, sklearn_clf, texts)
            out["topic"] = {
                "acc": round(accuracy(y_true, y_pred), 4),
                "macro_f1": round(macro_f1(y_true, y_pred, labels), 4),
            }
        if "retrieval" in args.tasks:
            if args.backend == "hf":
                q_emb = embed_texts(summaries, model_name=args.embed_model, batch_size=args.embed_batch_size)
                d_emb = embed_texts(texts, model_name=args.embed_model, batch_size=args.embed_batch_size)
                sim = q_emb @ d_emb.T
            else:
                sim = sklearn_retrieval_sim(sklearn_vectorizer, summaries, texts)
            out["summary_to_article_retrieval"] = {
                "recall@1": round(recall_at_k(sim, 1), 4),
                "recall@5": round(recall_at_k(sim, 5), 4),
                "recall@10": round(recall_at_k(sim, 10), 4),
            }
        return out

    # Original baseline
    print("Evaluating original...", flush=True)
    results["runs"]["original"] = eval_variant("original", originals)

    for rkeep in args.rkeeps:
        print(f"\n=== rkeep={rkeep} ===", flush=True)
        skeletons = [_cap_text(compressor(ex.article, rkeep), args.max_article_chars) for ex in test_examples]
        results["runs"].setdefault(str(rkeep), {})
        print("Evaluating skeleton...", flush=True)
        results["runs"][str(rkeep)]["skeleton"] = eval_variant("skeleton", skeletons)

        if reconstructor is not None:
            print("Reconstructing...", flush=True)
            rec_texts: list[str] = []
            for ex, sk in zip(test_examples, skeletons):
                rec = reconstructor.reconstruct(sk, original_len=len(ex.article))
                rec_texts.append(_cap_text(rec, args.max_article_chars))
            print("Evaluating reconstruction...", flush=True)
            results["runs"][str(rkeep)]["reconstructed"] = eval_variant("reconstructed", rec_texts)
        else:
            results["runs"][str(rkeep)]["reconstructed"] = None

    return results


def parse_args(argv: list[str]) -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--backend", type=str, default="sklearn", help="sklearn|hf")
    p.add_argument("--max-rows", type=int, default=None, help="Optional cap on rows loaded from the spreadsheet")
    p.add_argument("--n-train", type=int, default=1000, help="Train set size (sklearn backend only)")
    p.add_argument("--n-test", type=int, default=200)
    p.add_argument("--min-article-chars", type=int, default=200)
    p.add_argument("--max-article-chars", type=int, default=1500)
    p.add_argument("--max-summary-chars", type=int, default=400)

    p.add_argument("--compressor", type=str, default="wordlen", help="trunc|random|step|wordlen|entropy|wordfreq")
    p.add_argument("--rkeeps", type=float, nargs="+", default=[0.9, 0.7, 0.5, 0.3, 0.1])
    p.add_argument("--seed", type=int, default=0)

    p.add_argument("--decoder", type=str, default="none", help="none|gemini|hf")
    p.add_argument("--gemini-model", type=str, default="gemini-2.5-flash")  # default decoder; the published runs used gemini-2.0-flash, since retired
    p.add_argument("--gemini-api-key", type=str, default=None)
    p.add_argument("--hf-model", type=str, default="gpt2")

    p.add_argument("--tasks", type=str, nargs="+", default=["topic", "retrieval"], help="topic retrieval")
    # sklearn backend knobs
    p.add_argument("--tfidf-max-features", type=int, default=50000)
    p.add_argument("--tfidf-ngram-max", type=int, default=2)
    # hf backend knobs
    p.add_argument("--topic-model", type=str, default="facebook/bart-large-mnli")
    p.add_argument("--topic-batch-size", type=int, default=8)
    p.add_argument("--embed-model", type=str, default="sentence-transformers/all-MiniLM-L6-v2")
    p.add_argument("--embed-batch-size", type=int, default=32)

    p.add_argument("--out", type=str, default=str(RESULTS_PATH))
    return p.parse_args(argv)


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    results = run(args)
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w") as f:
        json.dump(results, f, indent=2)
    print(f"\nSaved results to {out_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
