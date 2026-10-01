#!/usr/bin/env python3
"""Compute BBC ROUGE-L statistics directly from workbook result sheets."""

from pathlib import Path
import statistics as st

import openpyxl
from rouge_score import rouge_scorer


DATA_PATH = Path(__file__).parent.parent / "data" / "BBC_news_latest.xlsx"

SPECS = {
    "Step": (
        "benchmark",
        "Articles",
        {
            "0.9": "polished_char_90",
            "0.8": "polished_char_80",
            "0.7": "polished_char_70",
            "0.6": "polished_char_60",
            "0.5": "polished_char_50",
            "0.4": "polished_char_40",
            "0.3": "polished_char_30",
            "0.2": "polished_char_20",
            "0.1": "polished_char_10",
        },
    ),
    "WordLen": (
        "adpative_word_length_deletion",
        "Articles",
        {
            "0.9": "polished_90",
            "0.8": "polished_80",
            "0.7": "polished_70",
            "0.6": "polished_60",
            "0.5": "polished_50",
            "0.4": "polished_40",
            "0.3": "polished_30",
            "0.2": "polished_20",
            "0.1": "polished_10",
        },
    ),
    "WordFreq": (
        "wf_3classes",
        "Articles",
        {
            "0.9": "polished_90",
            "0.8": "polished_80",
            "0.7": "polished_70",
            "0.6": "polished_60",
            "0.5": "polished_50",
            "0.4": "polished_40",
            "0.3": "polished_30",
            "0.2": "polished_20",
            "0.1": "polished_10",
        },
    ),
    "LP": (
        "wf_optimal",
        "Articles",
        {
            "0.9": "polished_optimal_keep_90",
            "0.8": "polished_optimal_keep_80",
            "0.7": "polished_optimal_keep_70",
            "0.6": "polished_optimal_keep_60",
            "0.5": "polished_optimal_keep_50",
            "0.4": "polished_optimal_keep_40",
            "0.3": "polished_optimal_keep_30",
            "0.2": "polished_optimal_keep_20",
            "0.1": "polished_optimal_keep_10",
        },
    ),
}


def main():
    scorer = rouge_scorer.RougeScorer(["rougeL"], use_stemmer=True)
    wb = openpyxl.load_workbook(DATA_PATH, data_only=True)
    for method, (sheet, ref_col, cols) in SPECS.items():
        ws = wb[sheet]
        headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        h = {name: i for i, name in enumerate(headers)}
        print(f"\n{method}")
        for rk, col in cols.items():
            vals = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                ref = row[h[ref_col]] if ref_col in h else None
                pred = row[h[col]] if col in h else None
                if not ref or not pred:
                    continue
                vals.append(scorer.score(str(ref), str(pred))["rougeL"].fmeasure)
            print(f"{rk}: mean={sum(vals)/len(vals):.4f} std={st.pstdev(vals):.4f} n={len(vals)}")
    wb.close()


if __name__ == "__main__":
    main()
