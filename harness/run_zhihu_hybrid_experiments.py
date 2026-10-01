#!/usr/bin/env python3
"""
Entropy and hybrid deletion experiments on the Zhihu dataset.

This mirrors the English BBC/Wikipedia hybrid scripts, but operates on the
segmented Zhihu test split and uses Chinese models throughout:

  - Decoder: Gemini 2.0 Flash in the published runs (since retired); the default is now gemini-2.5-flash
  - Surprisal model: uer/gpt2-chinese-cluecorpussmall
  - Evaluation: BERTScore with bert-base-chinese

Methods:
  1. Entropy (pure surprisal-based deletion)
  2. Entropy-LP
  3. Entropy-in-FreqBuckets
  4. Hybrid-alpha for alpha in {0.3, 0.5, 0.7}

Results are saved incrementally to zhihu_hybrid_results.json so interrupted
runs can resume.
"""

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

sys.stdout.reconfigure(line_buffering=True)
sys.stderr.reconfigure(line_buffering=True)

import numpy as np
import openpyxl
import torch
from bert_score import score as bert_score_fn
from transformers import AutoModelForCausalLM, AutoTokenizer
from wordfreq import freq_to_zipf, get_frequency_dict, smash_numbers


DATA_PATH = Path(__file__).parent.parent / "data" / "Zhihu_Dataset.xlsx"
RESULTS_PATH = Path(__file__).parent / "zhihu_hybrid_results.json"
DEFAULT_RATES = [0.9, 0.7, 0.5, 0.3, 0.1]
DEFAULT_N_TEST = 200
DEFAULT_N_CALIB = 50
ALPHA_VALUES = [0.3, 0.5, 0.7]

FREQ_HIGH = 5.0
FREQ_MID = 3.0
ZH_GPT_MODEL = "uer/gpt2-chinese-cluecorpussmall"
ZH_BERTSCORE_MODEL = "bert-base-chinese"
ZH_FREQS = get_frequency_dict("zh", "best")

from dotenv import load_dotenv

load_dotenv(os.path.expanduser("~/.env"))
GOOGLE_API_KEY = os.environ["GOOGLE_API_KEY"]  # set your own key
if not GOOGLE_API_KEY:
    sys.exit("GOOGLE_API_KEY not found")

import google.generativeai as genai

genai.configure(api_key=GOOGLE_API_KEY)
gemini = genai.GenerativeModel("gemini-2.5-flash")  # default decoder; the published runs used gemini-2.0-flash, since retired


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--n-test", type=int, default=DEFAULT_N_TEST)
    parser.add_argument("--n-calib", type=int, default=DEFAULT_N_CALIB)
    parser.add_argument(
        "--rates",
        type=float,
        nargs="+",
        default=DEFAULT_RATES,
        help="Retention rates to evaluate, e.g. --rates 0.9 0.7 0.5 0.3 0.1",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Load data/models and run one local deletion pass without Gemini calls.",
    )
    return parser.parse_args()


def load_results():
    if RESULTS_PATH.exists():
        with open(RESULTS_PATH) as f:
            return json.load(f)
    return {}


def save_results(all_results):
    with open(RESULTS_PATH, "w") as f:
        json.dump(all_results, f, indent=2, ensure_ascii=False)


def split_segmented(segmented):
    if not segmented:
        return []
    return [tok for tok in str(segmented).split("/") if tok]


def normalize_compact_text(text):
    return re.sub(r"\s+", "", str(text))


def load_examples(n):
    wb = openpyxl.load_workbook(str(DATA_PATH), read_only=True)
    ws = wb["Testing(Word Segmentation)"]

    headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    h = {name: idx for idx, name in enumerate(headers)}

    examples = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_text = str(row[h["RESPONSE"]]).strip() if row[h["RESPONSE"]] else ""
        segmented = row[h["tok_flat_coarse"]] if "tok_flat_coarse" in h else None
        tokens = split_segmented(segmented)
        compact_text = "".join(tokens) if tokens else normalize_compact_text(raw_text)
        if not raw_text or not compact_text or len(compact_text) < 80:
            continue
        examples.append(
            {
                "raw_text": raw_text,
                "compact_text": compact_text,
                "tokens": tokens if tokens else list(compact_text),
            }
        )
        if len(examples) >= n:
            break

    wb.close()
    return examples


def call_gemini(prompt, max_retries=3):
    for attempt in range(max_retries):
        try:
            response = gemini.generate_content(prompt)
            return response.text.strip()
        except Exception as e:
            print(f"  Gemini attempt {attempt + 1} failed: {e}", flush=True)
            if attempt < max_retries - 1:
                time.sleep(2)
            else:
                return ""


def reconstruct_prompt(degraded_text, original_len):
    return (
        "你是一个中文文本重建助手。输入是一段知乎回答，其中部分词语已被删除。"
        "请在不添加输入中没有依据的新事实、新人物或新事件的前提下，尽可能忠实地重建原文。"
        f"输出长度应大致接近原文（约{original_len}个字符）。"
        "只返回重建后的中文文本，不要加解释。\n\n"
        f"残缺文本：\n{degraded_text}"
    )


def load_zh_gpt():
    print(f"Loading Chinese GPT model: {ZH_GPT_MODEL}", flush=True)
    tokenizer = AutoTokenizer.from_pretrained(ZH_GPT_MODEL)
    model = AutoModelForCausalLM.from_pretrained(ZH_GPT_MODEL)
    model.eval()
    if torch.cuda.is_available():
        model = model.cuda()
    return tokenizer, model


def compute_char_surprisal(text, tokenizer, model):
    encoded = tokenizer(
        text,
        return_tensors="pt",
        return_offsets_mapping=True,
        truncation=True,
        max_length=1024,
        add_special_tokens=False,
    )
    offsets = encoded.pop("offset_mapping")[0].tolist()
    if torch.cuda.is_available():
        encoded = {k: v.cuda() for k, v in encoded.items()}

    with torch.no_grad():
        outputs = model(**encoded)
        logits = outputs.logits

    shift_logits = logits[:, :-1, :]
    shift_labels = encoded["input_ids"][:, 1:]
    log_probs = torch.nn.functional.log_softmax(shift_logits, dim=-1)
    token_log_probs = log_probs.gather(2, shift_labels.unsqueeze(-1)).squeeze(-1)
    surprisal = -token_log_probs.squeeze(0).cpu().numpy()

    if len(surprisal) == 0:
        return np.array([0.0]), [(0, len(text))]

    median_s = float(np.median(surprisal))
    surprisal = np.concatenate([[median_s], surprisal])
    return surprisal, offsets


def compute_word_surprisal(compact_text, words, tokenizer, model):
    token_surprisal, offsets = compute_char_surprisal(compact_text, tokenizer, model)

    spans = []
    pos = 0
    for word in words:
        end = pos + len(word)
        spans.append((pos, end))
        pos = end

    if pos != len(compact_text):
        compact_text = "".join(words)
        token_surprisal, offsets = compute_char_surprisal(compact_text, tokenizer, model)
        spans = []
        pos = 0
        for word in words:
            end = pos + len(word)
            spans.append((pos, end))
            pos = end

    word_scores = []
    for start, end in spans:
        vals = []
        for s, (tok_start, tok_end) in zip(token_surprisal, offsets):
            if tok_end <= start or tok_start >= end:
                continue
            vals.append(float(s))
        if vals:
            word_scores.append(float(np.mean(vals)))
        else:
            word_scores.append(float(np.median(token_surprisal)))
    return np.array(word_scores, dtype=float)


def get_zipf_scores(words):
    scores = []
    for word in words:
        token = str(word).strip()
        if not token:
            scores.append(0.0)
            continue
        if re.fullmatch(r"\W+", token):
            scores.append(0.0)
            continue
        smashed = smash_numbers(token)
        freq = ZH_FREQS.get(smashed)
        if freq is None:
            char_freqs = [ZH_FREQS.get(ch) for ch in token if ZH_FREQS.get(ch) is not None]
            if char_freqs:
                freq = float(np.mean(char_freqs))
            else:
                freq = 1e-9
        scores.append(round(freq_to_zipf(freq), 2))
    return np.array(scores, dtype=float)


def assign_freq_bucket(zipf_scores):
    buckets = np.zeros(len(zipf_scores), dtype=int)
    buckets[zipf_scores >= FREQ_HIGH] = 2
    buckets[(zipf_scores >= FREQ_MID) & (zipf_scores < FREQ_HIGH)] = 1
    return buckets


def rank_normalize(arr):
    n = len(arr)
    if n <= 1:
        return np.zeros(n)
    order = np.argsort(arr, kind="stable")
    ranks = np.empty(n)
    ranks[order] = np.arange(n)
    return ranks / (n - 1)


def solve_lp(b_full, bucket_sizes, n_delete_total):
    cost = 1.0 - b_full
    order = np.argsort(cost)
    w = np.zeros(len(b_full))
    remaining = n_delete_total
    for k in order:
        cap = int(bucket_sizes[k])
        if cap <= 0:
            continue
        if remaining >= cap:
            w[k] = 1.0
            remaining -= cap
        else:
            w[k] = remaining / cap
            remaining = 0
            break
    return w


def reconstruct_text(words, keep_mask):
    kept = [word for word, keep in zip(words, keep_mask) if keep]
    return "".join(kept)


def evaluate_reconstructions(preds, refs):
    _, _, f1 = bert_score_fn(
        preds,
        refs,
        model_type=ZH_BERTSCORE_MODEL,
        lang="zh",
        verbose=False,
        device="cpu",
    )
    return f1.mean().item(), f1.std().item()


def calibrate_surprisal_buckets(calib_examples, tokenizer, model, all_results):
    if "calib_surprisal" in all_results:
        d = all_results["calib_surprisal"]
        print("  Surprisal calibration already exists, skipping.", flush=True)
        return np.array(d["b_full"]), d["thresholds"]

    print(f"\n--- Calibrating surprisal buckets (n={len(calib_examples)}) ---", flush=True)
    all_scores = []
    for ex in calib_examples:
        all_scores.extend(compute_word_surprisal(ex["compact_text"], ex["tokens"], tokenizer, model).tolist())
    t33 = float(np.percentile(all_scores, 33.33))
    t67 = float(np.percentile(all_scores, 66.67))
    print(f"  Surprisal thresholds: 33%={t33:.3f}, 67%={t67:.3f}", flush=True)

    bucket_recs = [[], [], []]
    bucket_refs = [[], [], []]
    for i, ex in enumerate(calib_examples):
        words = ex["tokens"]
        surp = compute_word_surprisal(ex["compact_text"], words, tokenizer, model)
        for bk in range(3):
            if bk == 0:
                delete_mask = surp <= t33
            elif bk == 1:
                delete_mask = (surp > t33) & (surp <= t67)
            else:
                delete_mask = surp > t67
            degraded = reconstruct_text(words, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            bucket_recs[bk].append(rec)
            bucket_refs[bk].append(ex["raw_text"])
            time.sleep(0.1)
        if (i + 1) % 10 == 0:
            print(f"  calibration {i + 1}/{len(calib_examples)}", flush=True)

    b_full = np.zeros(3)
    for bk in range(3):
        mean_f1, _ = evaluate_reconstructions(bucket_recs[bk], bucket_refs[bk])
        b_full[bk] = mean_f1
        print(f"  Surprisal bucket {bk}: B_k^full={b_full[bk]:.4f}", flush=True)

    all_results["calib_surprisal"] = {"b_full": b_full.tolist(), "thresholds": [t33, t67]}
    save_results(all_results)
    return b_full, [t33, t67]


def calibrate_freq_buckets(calib_examples, tokenizer, model, all_results):
    if "calib_frequency" in all_results:
        d = all_results["calib_frequency"]
        print("  Frequency calibration already exists, skipping.", flush=True)
        return np.array(d["b_full"])

    print(f"\n--- Calibrating frequency buckets (n={len(calib_examples)}) ---", flush=True)
    bucket_recs = [[], [], []]
    bucket_refs = [[], [], []]
    for i, ex in enumerate(calib_examples):
        words = ex["tokens"]
        zipf_scores = get_zipf_scores(words)
        bucket_ids = assign_freq_bucket(zipf_scores)
        for bk in range(3):
            delete_mask = bucket_ids == bk
            degraded = reconstruct_text(words, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            bucket_recs[bk].append(rec)
            bucket_refs[bk].append(ex["raw_text"])
            time.sleep(0.1)
        if (i + 1) % 10 == 0:
            print(f"  calibration {i + 1}/{len(calib_examples)}", flush=True)

    b_full = np.zeros(3)
    for bk in range(3):
        mean_f1, _ = evaluate_reconstructions(bucket_recs[bk], bucket_refs[bk])
        b_full[bk] = mean_f1
        print(f"  Frequency bucket {bk}: B_k^full={b_full[bk]:.4f}", flush=True)

    all_results["calib_frequency"] = {"b_full": b_full.tolist()}
    save_results(all_results)
    return b_full


def run_entropy(examples, tokenizer, model, retention_rates, all_results):
    results = all_results.get("entropy", {})
    for rkeep in retention_rates:
        key = str(rkeep)
        if key in results:
            print(f"  entropy rkeep={rkeep} already done, skipping.", flush=True)
            continue
        print(f"\n=== Entropy rkeep={rkeep} ===", flush=True)
        preds, refs = [], []
        for i, ex in enumerate(examples):
            words = ex["tokens"]
            surp = compute_word_surprisal(ex["compact_text"], words, tokenizer, model)
            n_keep = max(1, int(len(words) * rkeep))
            keep_idx = set(np.argsort(surp)[-n_keep:])
            keep_mask = np.array([j in keep_idx for j in range(len(words))], dtype=bool)
            degraded = reconstruct_text(words, keep_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            preds.append(rec)
            refs.append(ex["raw_text"])
            if (i + 1) % 10 == 0:
                print(f"  {i + 1}/{len(examples)} done", flush=True)
            time.sleep(0.1)

        mean_f1, std_f1 = evaluate_reconstructions(preds, refs)
        print(f"  BERTScore F1: {mean_f1:.4f} +/- {std_f1:.4f}", flush=True)
        results[key] = {"mean_f1": round(mean_f1, 4), "std_f1": round(std_f1, 4), "n": len(examples)}
        all_results["entropy"] = results
        save_results(all_results)
    return results


def run_entropy_lp(examples, tokenizer, model, b_full, thresholds, retention_rates, all_results):
    results = all_results.get("entropy_lp", {})
    t33, t67 = thresholds
    for rkeep in retention_rates:
        key = str(rkeep)
        if key in results:
            print(f"  entropy-lp rkeep={rkeep} already done, skipping.", flush=True)
            continue
        print(f"\n=== Entropy-LP rkeep={rkeep} ===", flush=True)
        preds, refs = [], []
        for i, ex in enumerate(examples):
            words = ex["tokens"]
            surp = compute_word_surprisal(ex["compact_text"], words, tokenizer, model)
            n = len(words)
            n_del = max(0, int(n * (1 - rkeep)))
            bucket_ids = np.zeros(n, dtype=int)
            bucket_ids[(surp > t33) & (surp <= t67)] = 1
            bucket_ids[surp > t67] = 2
            bucket_sizes = np.array([np.sum(bucket_ids == k) for k in range(3)])
            w = solve_lp(b_full, bucket_sizes, n_del)

            delete_mask = np.zeros(n, dtype=bool)
            for bk in range(3):
                idx = np.where(bucket_ids == bk)[0]
                n_del_bk = int(round(w[bk] * len(idx)))
                if n_del_bk > 0:
                    cheapest = idx[np.argsort(surp[idx])][:n_del_bk]
                    delete_mask[cheapest] = True

            degraded = reconstruct_text(words, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            preds.append(rec)
            refs.append(ex["raw_text"])
            if (i + 1) % 10 == 0:
                print(f"  {i + 1}/{len(examples)} done", flush=True)
            time.sleep(0.1)

        mean_f1, std_f1 = evaluate_reconstructions(preds, refs)
        print(f"  BERTScore F1: {mean_f1:.4f} +/- {std_f1:.4f}", flush=True)
        results[key] = {"mean_f1": round(mean_f1, 4), "std_f1": round(std_f1, 4), "n": len(examples)}
        all_results["entropy_lp"] = results
        save_results(all_results)
    return results


def run_entropy_in_freqbuckets(examples, tokenizer, model, b_full, retention_rates, all_results):
    results = all_results.get("entropy_in_freqbuckets", {})
    for rkeep in retention_rates:
        key = str(rkeep)
        if key in results:
            print(f"  entropy-in-freqbuckets rkeep={rkeep} already done, skipping.", flush=True)
            continue
        print(f"\n=== Entropy-in-FreqBuckets rkeep={rkeep} ===", flush=True)
        preds, refs = [], []
        for i, ex in enumerate(examples):
            words = ex["tokens"]
            surp = compute_word_surprisal(ex["compact_text"], words, tokenizer, model)
            zipf_scores = get_zipf_scores(words)
            n = len(words)
            n_del = max(0, int(n * (1 - rkeep)))
            bucket_ids = assign_freq_bucket(zipf_scores)
            bucket_sizes = np.array([np.sum(bucket_ids == k) for k in range(3)])
            w = solve_lp(b_full, bucket_sizes, n_del)

            delete_mask = np.zeros(n, dtype=bool)
            for bk in range(3):
                idx = np.where(bucket_ids == bk)[0]
                n_del_bk = int(round(w[bk] * len(idx)))
                if n_del_bk > 0:
                    cheapest = idx[np.argsort(surp[idx])][:n_del_bk]
                    delete_mask[cheapest] = True

            degraded = reconstruct_text(words, ~delete_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            preds.append(rec)
            refs.append(ex["raw_text"])
            if (i + 1) % 10 == 0:
                print(f"  {i + 1}/{len(examples)} done", flush=True)
            time.sleep(0.1)

        mean_f1, std_f1 = evaluate_reconstructions(preds, refs)
        print(f"  BERTScore F1: {mean_f1:.4f} +/- {std_f1:.4f}", flush=True)
        results[key] = {"mean_f1": round(mean_f1, 4), "std_f1": round(std_f1, 4), "n": len(examples)}
        all_results["entropy_in_freqbuckets"] = results
        save_results(all_results)
    return results


def run_hybrid_alpha(examples, tokenizer, model, alpha, retention_rates, all_results):
    result_key = f"hybrid_alpha_{alpha}"
    results = all_results.get(result_key, {})
    for rkeep in retention_rates:
        key = str(rkeep)
        if key in results:
            print(f"  alpha={alpha} rkeep={rkeep} already done, skipping.", flush=True)
            continue
        print(f"\n=== Hybrid-alpha={alpha} rkeep={rkeep} ===", flush=True)
        preds, refs = [], []
        for i, ex in enumerate(examples):
            words = ex["tokens"]
            surp = compute_word_surprisal(ex["compact_text"], words, tokenizer, model)
            zipf_scores = get_zipf_scores(words)
            n_keep = max(1, int(len(words) * rkeep))

            rank_freq = rank_normalize(zipf_scores)
            inv_surp = np.where(surp > 0, 1.0 / surp, 0.0)
            rank_pred = rank_normalize(inv_surp)
            combined = alpha * rank_freq + (1.0 - alpha) * rank_pred
            keep_idx = set(np.argsort(combined)[:n_keep])
            keep_mask = np.array([j in keep_idx for j in range(len(words))], dtype=bool)

            degraded = reconstruct_text(words, keep_mask)
            rec = call_gemini(reconstruct_prompt(degraded, len(ex["raw_text"])))
            preds.append(rec)
            refs.append(ex["raw_text"])
            if (i + 1) % 10 == 0:
                print(f"  {i + 1}/{len(examples)} done", flush=True)
            time.sleep(0.1)

        mean_f1, std_f1 = evaluate_reconstructions(preds, refs)
        print(f"  BERTScore F1: {mean_f1:.4f} +/- {std_f1:.4f}", flush=True)
        results[key] = {"mean_f1": round(mean_f1, 4), "std_f1": round(std_f1, 4), "n": len(examples)}
        all_results[result_key] = results
        save_results(all_results)
    return results


def print_summary(all_results, retention_rates):
    print("\n" + "=" * 100, flush=True)
    print("SUMMARY (Zhihu BERTScore F1)", flush=True)
    print("=" * 100, flush=True)
    cols = [
        ("Entropy", all_results.get("entropy", {})),
        ("Entropy-LP", all_results.get("entropy_lp", {})),
        ("Ent.-FreqBkt", all_results.get("entropy_in_freqbuckets", {})),
        ("Hybrid-0.3", all_results.get("hybrid_alpha_0.3", {})),
        ("Hybrid-0.5", all_results.get("hybrid_alpha_0.5", {})),
        ("Hybrid-0.7", all_results.get("hybrid_alpha_0.7", {})),
    ]
    header = f"{'rkeep':<8}" + "".join(f"{label:<16}" for label, _ in cols)
    print(header, flush=True)
    print("-" * len(header), flush=True)
    for rkeep in retention_rates:
        key = str(rkeep)
        row = f"{rkeep:<8}"
        for _, data in cols:
            val = data.get(key, {}).get("mean_f1", "---") if isinstance(data, dict) else "---"
            row += f"{str(val):<16}"
        print(row, flush=True)


def run_dry_run(example, tokenizer, model):
    words = example["tokens"]
    surp = compute_word_surprisal(example["compact_text"], words, tokenizer, model)
    zipf_scores = get_zipf_scores(words)
    keep_idx = set(np.argsort(surp)[-max(1, int(len(words) * 0.5)):])
    degraded = reconstruct_text(words, np.array([i in keep_idx for i in range(len(words))], dtype=bool))
    print("Dry run OK.", flush=True)
    print(f"Words: {len(words)}", flush=True)
    print(f"Mean surprisal: {float(np.mean(surp)):.4f}", flush=True)
    print(f"Mean zh zipf: {float(np.mean(zipf_scores)):.4f}", flush=True)
    print(f"Degraded preview: {degraded[:180]}", flush=True)


def main():
    args = parse_args()
    print("Loading Zhihu examples...", flush=True)
    examples = load_examples(args.n_test)
    print(f"Loaded {len(examples)} examples from Testing(Word Segmentation)", flush=True)
    print(f"Mean compact length: {np.mean([len(ex['compact_text']) for ex in examples]):.1f} chars", flush=True)

    all_results = load_results()
    if all_results:
        print(f"Loaded existing partial results from {RESULTS_PATH}", flush=True)

    tokenizer, model = load_zh_gpt()

    if args.dry_run:
        run_dry_run(examples[0], tokenizer, model)
        return

    calib_examples = examples[: args.n_calib]
    retention_rates = args.rates

    print("\n" + "=" * 60, flush=True)
    print("CALIBRATION", flush=True)
    print("=" * 60, flush=True)
    b_full_surp, thresholds = calibrate_surprisal_buckets(calib_examples, tokenizer, model, all_results)
    b_full_freq = calibrate_freq_buckets(calib_examples, tokenizer, model, all_results)

    print("\n" + "=" * 60, flush=True)
    print("ENTROPY", flush=True)
    print("=" * 60, flush=True)
    run_entropy(examples, tokenizer, model, retention_rates, all_results)

    print("\n" + "=" * 60, flush=True)
    print("ENTROPY-LP", flush=True)
    print("=" * 60, flush=True)
    run_entropy_lp(examples, tokenizer, model, b_full_surp, thresholds, retention_rates, all_results)

    print("\n" + "=" * 60, flush=True)
    print("ENTROPY-IN-FREQBUCKETS", flush=True)
    print("=" * 60, flush=True)
    run_entropy_in_freqbuckets(examples, tokenizer, model, b_full_freq, retention_rates, all_results)

    print("\n" + "=" * 60, flush=True)
    print("HYBRID-ALPHA", flush=True)
    print("=" * 60, flush=True)
    for alpha in ALPHA_VALUES:
        run_hybrid_alpha(examples, tokenizer, model, alpha, retention_rates, all_results)

    save_results(all_results)
    print(f"\nAll results saved to {RESULTS_PATH}", flush=True)
    print_summary(all_results, retention_rates)


if __name__ == "__main__":
    main()
