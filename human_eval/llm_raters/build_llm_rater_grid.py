#!/usr/bin/env python3
"""Build the LLM-rater grid: items, facts, blind batches, item-to-condition mapping.

Grid: 100 BBC articles x {Step, WordLen, WordFreq (= WordFreq-3), LP} x 9 rates (workbook, gemini-2.0-flash decoder)
    + 100 BBC chunks x {Entropy, Entropy-LP, Hybrid-0.5} x 5 rates (samples/bbc_hybrid_samples.jsonl, gemini-2.5-flash).
Batch 000 = the 54 items of the human study (calibration anchor), verbatim.

Facts are drafted from the ORIGINAL text only with gemini-2.5-flash (same drafting prompt as the human study),
cached in facts.json so re-runs never re-draft; the human study's verified facts are reused where the original
text matches.

Usage: python build_llm_rater_grid.py [--workbook PATH] [--articles 100] [--chunks 100] [--no-draft]
"""
import argparse, ast, json, os, random, re, sys, time, hashlib
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
HUMAN_DIR = REPO / "human_eval"

WORD_SPECS = {  # strategy -> (sheet, original col, {rate: recon col}), as in harness/compute_bbc_rougel.py
    "Step":     ("benchmark", "Articles", {r: f"polished_char_{r}" for r in (90, 80, 70, 60, 50, 40, 30, 20, 10)}),
    "WordLen":  ("adpative_word_length_deletion", "Articles", {r: f"polished_{r}" for r in (90, 80, 70, 60, 50, 40, 30, 20, 10)}),
    "WordFreq": ("wf_3classes", "Articles", {r: f"polished_{r}" for r in (90, 80, 70, 60, 50, 40, 30, 20, 10)}),
    "LP":       ("wf_optimal", "Articles", {r: f"polished_optimal_keep_{r}" for r in (90, 80, 70, 60, 50, 40, 30, 20, 10)}),
}
SEM_METHODS = {"entropy": "Entropy", "entropy_lp": "Entropy-LP", "hybrid_alpha_0.5": "Hybrid-0.5"}

FACT_PROMPT = """Extract 3 to 5 atomic factual statements from the news excerpt below.

Rules:
- Each fact must be a single, self-contained claim (one subject, one predicate).
- Prefer facts anchored on named entities, numbers, dates, places, or quantities.
- Use only information stated in the excerpt; no outside knowledge, no inference.
- Phrase each fact as a short declarative sentence.
- Return ONLY a JSON array of strings, nothing else.

Excerpt:
\"\"\"{text}\"\"\"
"""


def load_env():
    p = Path.home() / ".env"
    if p.exists():
        for line in p.read_text().splitlines():
            if "=" in line and not line.strip().startswith("#"):
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def norm(t):
    return re.sub(r"\s+", " ", t or "").strip()


def key_of(text):
    return hashlib.sha1(norm(text).encode()).hexdigest()[:16]


def load_workbook_items(path, n_articles):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = {}
    for strat, (sheet, ocol, cols) in WORD_SPECS.items():
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        hdr = [str(h) for h in next(rows)]
        idx = {h: i for i, h in enumerate(hdr)}
        recs = {}
        for row in rows:
            o = row[idx[ocol]] if idx[ocol] < len(row) else None
            if not isinstance(o, str) or len(o) < 50:
                continue
            recs[key_of(o)] = (o, {r: (row[idx[c]] if c in idx and idx[c] < len(row) else None) for r, c in cols.items()})
        sheets[strat] = recs
        print(f"  sheet {sheet}: {len(recs)} articles", file=sys.stderr)
    common = [k for k in sheets["Step"] if all(k in sheets[s] for s in WORD_SPECS)]
    print(f"  articles present in all four sheets: {len(common)}", file=sys.stderr)
    rng = random.Random(20260913)
    rng.shuffle(common)
    chosen = common[:n_articles]
    items = []
    for k in chosen:
        orig = sheets["Step"][k][0]
        for strat in WORD_SPECS:
            for r, rec in sheets[strat][k][1].items():
                if isinstance(rec, str) and rec.strip():
                    items.append(dict(unit="article", unit_key=k, strategy=strat, rkeep=r / 100,
                                      decoder="gemini-2.0-flash", original=orig, reconstruction=rec))
    return items


def load_semantic_items(path, n_chunks):
    rows = [json.loads(l) for l in open(path)]
    rows = [x for x in rows if x["method"] in SEM_METHODS]
    by_chunk = {}
    for x in rows:
        by_chunk.setdefault(x["sample_idx"], []).append(x)
    rng = random.Random(20260913)
    idxs = sorted(by_chunk)
    rng.shuffle(idxs)
    items = []
    for si in idxs[:n_chunks]:
        for x in by_chunk[si]:
            if isinstance(x["reconstruction"], str) and x["reconstruction"].strip():
                items.append(dict(unit="chunk", unit_key=f"chunk{si}", strategy=SEM_METHODS[x["method"]], rkeep=float(x["rkeep"]),
                                  decoder=x["decoder_model"], original=x["reference"], reconstruction=x["reconstruction"]))
    print(f"  semantic items: {len(items)} from {min(n_chunks, len(idxs))} chunks", file=sys.stderr)
    return items


def draft_facts(originals, cache_path, do_draft=True):
    cache = json.load(open(cache_path)) if cache_path.exists() else {}
    # reuse the human study's verified facts where the original text matches
    e5_items = json.load(open(HUMAN_DIR / "e5_items.json"))
    e5_facts = json.load(open(HUMAN_DIR / "e5_facts.json"))
    for it in e5_items:
        k = key_of(it["original"])
        if k not in cache and str(it["id"]) in e5_facts:
            cache[k] = {"facts": e5_facts[str(it["id"])], "source": "e5_human_verified"}
    todo = [(k, t) for k, t in originals.items() if k not in cache]
    print(f"  facts cached: {len(cache)}, to draft: {len(todo)}", file=sys.stderr)
    if todo and do_draft:
        from google import genai
        from google.genai import types
        client = genai.Client(api_key=os.environ["GOOGLE_API_KEY"], http_options=types.HttpOptions(timeout=60000))
        for i, (k, text) in enumerate(todo, 1):
            for attempt in range(5):
                try:
                    resp = client.models.generate_content(model="gemini-2.5-flash", contents=FACT_PROMPT.format(text=text))
                    raw = resp.text.strip()
                    raw = re.sub(r"^```(?:json)?|```$", "", raw, flags=re.M).strip()
                    facts = json.loads(raw)
                    assert isinstance(facts, list) and 2 <= len(facts) <= 6 and all(isinstance(f, str) for f in facts)
                    cache[k] = {"facts": facts[:5], "source": "gemini-2.5-flash"}
                    break
                except Exception as e:
                    print(f"    retry {attempt+1} on {k}: {e}", file=sys.stderr)
                    time.sleep(2 + 3 * attempt)
            if i % 10 == 0 or i == len(todo):
                json.dump(cache, open(cache_path, "w"), indent=1, ensure_ascii=False)
                print(f"    drafted {i}/{len(todo)}", file=sys.stderr)
    json.dump(cache, open(cache_path, "w"), indent=1, ensure_ascii=False)
    return cache


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default=str(REPO / "data" / "BBC_news_latest.xlsx"))
    ap.add_argument("--hybrid", default=str(REPO / "samples" / "bbc_hybrid_samples.jsonl"))
    ap.add_argument("--articles", type=int, default=100)
    ap.add_argument("--chunks", type=int, default=100)
    ap.add_argument("--batch-size", type=int, default=25)
    ap.add_argument("--no-draft", action="store_true")
    args = ap.parse_args()
    load_env()

    print("Loading items...", file=sys.stderr)
    items = load_workbook_items(args.workbook, args.articles) + load_semantic_items(args.hybrid, args.chunks)
    originals = {}
    for it in items:
        originals.setdefault(key_of(it["original"]), it["original"])
    print(f"  grid items: {len(items)}; distinct originals: {len(originals)}", file=sys.stderr)

    facts = draft_facts(originals, HERE / "facts.json", do_draft=not args.no_draft)
    missing = [k for k in originals if k not in facts]
    if missing:
        print(f"!! {len(missing)} originals still lack facts; run again without --no-draft", file=sys.stderr)

    # batch 000: the 54 human-study items verbatim (calibration anchor)
    e5_items = json.load(open(HUMAN_DIR / "e5_items.json"))
    e5_facts = json.load(open(HUMAN_DIR / "e5_facts.json"))
    batches = {}
    mapping = {}
    b0 = []
    for it in e5_items:
        iid = f"e5_{int(it['id']):03d}"
        b0.append({"item_id": iid, "original": it["original"], "reconstruction": it["reconstruction"], "facts": e5_facts[str(it["id"])]})
        mapping[iid] = {"set": "e5_calibration", "e5_id": it["id"]}
    batches[0] = b0

    rng = random.Random(20260913)
    order = list(range(len(items)))
    rng.shuffle(order)
    grid = []
    for n, i in enumerate(order):
        it = items[i]
        k = key_of(it["original"])
        if k not in facts:
            continue
        iid = f"g_{n:05d}"
        grid.append({"item_id": iid, "original": it["original"], "reconstruction": it["reconstruction"], "facts": facts[k]["facts"]})
        mapping[iid] = {"set": "grid", "unit": it["unit"], "unit_key": it["unit_key"], "strategy": it["strategy"],
                        "rkeep": it["rkeep"], "decoder": it["decoder"], "facts_source": facts[k]["source"]}
    for b, start in enumerate(range(0, len(grid), args.batch_size), start=1):
        batches[b] = grid[start:start + args.batch_size]

    (HERE / "batches").mkdir(exist_ok=True)
    for b, its in batches.items():
        json.dump(its, open(HERE / "batches" / f"batch_{b:03d}.json", "w"), indent=1, ensure_ascii=False)
    json.dump(mapping, open(HERE / "mapping.json", "w"), indent=1, ensure_ascii=False)
    from collections import Counter
    c = Counter((m["strategy"], m["rkeep"]) for m in mapping.values() if m["set"] == "grid")
    summary = {"n_batches": len(batches), "n_grid_items": len(grid), "n_calibration_items": len(b0),
               "cells": {f"{s}@{r}": n for (s, r), n in sorted(c.items())}}
    json.dump(summary, open(HERE / "build_summary.json", "w"), indent=1)
    print(json.dumps(summary, indent=1))


if __name__ == "__main__":
    main()
